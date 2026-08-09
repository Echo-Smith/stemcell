from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from build_deepseek_resume_docx import (
    ACCENT_DEEP,
    BODY,
    INK,
    MUTED,
    add_body,
    add_custom_bullet_numbering,
    add_first_page_avatar,
    add_hyperlink,
    add_item_heading,
    add_meta,
    add_section_heading,
    apply_bullet,
    configure_styles,
    set_cell_free_font,
    set_paragraph_rhythm,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docx" / "莫思意-AI产品经理-定向简历.docx"


def add_page_fields(paragraph):
    paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(186), WD_TAB_ALIGNMENT.RIGHT)
    left = paragraph.add_run("莫思意 | AI 产品经理")
    set_cell_free_font(left, 7.0, MUTED)
    paragraph.add_run("\t")
    for field_name in ["PAGE", "NUMPAGES"]:
        if field_name == "NUMPAGES":
            slash = paragraph.add_run(" / ")
            set_cell_free_font(slash, 7.0, MUTED)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), field_name)
        run = OxmlElement("w:r")
        prop = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), MUTED)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "14")
        prop.extend([color, size])
        text = OxmlElement("w:t")
        text.text = "1"
        run.extend([prop, text])
        field.append(run)
        paragraph._p.append(field)


def add_compact_bullet(doc, num_id, label, text, size=8.0):
    paragraph = doc.add_paragraph()
    apply_bullet(paragraph, num_id)
    set_paragraph_rhythm(paragraph, after=0.6, line=1.04, keep=True)
    if label:
        set_cell_free_font(paragraph.add_run(label), size, ACCENT_DEEP, bold=True)
    set_cell_free_font(paragraph.add_run(text), size, BODY)
    return paragraph


def add_meta_links(doc, text, links):
    paragraph = doc.add_paragraph(style="Resume Meta")
    set_paragraph_rhythm(paragraph, after=0.8, line=1.0, keep=True)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(190), WD_TAB_ALIGNMENT.RIGHT)
    set_cell_free_font(paragraph.add_run(text), 6.9, MUTED)
    paragraph.add_run("\t")
    for index, (label, url) in enumerate(links):
        if index:
            set_cell_free_font(paragraph.add_run(" | "), 6.7, MUTED)
        add_hyperlink(paragraph, label, url, 6.8)
    return paragraph


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(7)
    section.bottom_margin = Mm(6)
    section.left_margin = Mm(9)
    section.right_margin = Mm(9)
    section.header_distance = Mm(3)
    section.footer_distance = Mm(3.5)
    add_first_page_avatar(section)

    configure_styles(doc)
    doc.styles["Normal"].font.size = Pt(8.2)
    doc.styles["Normal"].paragraph_format.space_after = Pt(1.4)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.04
    doc.styles["Heading 1"].font.size = Pt(9.6)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(3.2)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(2.2)
    doc.styles["Heading 2"].font.size = Pt(8.9)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(2.2)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(0.6)
    bullet_id = add_custom_bullet_numbering(doc)

    title = doc.add_paragraph()
    set_paragraph_rhythm(title, after=1.0, line=1.0)
    set_cell_free_font(title.add_run("莫思意"), 19.5, INK, bold=True)

    target = doc.add_paragraph()
    set_paragraph_rhythm(target, after=1.5, line=1.0)
    set_cell_free_font(target.add_run("AI 产品经理｜AI 应用 / Agent 工作流"), 9.4, ACCENT_DEEP, bold=True)

    contact = doc.add_paragraph()
    set_paragraph_rhythm(contact, after=2.0, line=1.0)
    for value in ["19290865729", "aodomosiyi@foxmail.com", "杭州"]:
        set_cell_free_font(contact.add_run(value + "  |  "), 7.0, MUTED)
    add_hyperlink(contact, "作品集", "https://profile.ericdocmic.top", 7.0)
    set_cell_free_font(contact.add_run("  |  "), 7.0, MUTED)
    add_hyperlink(contact, "GitHub", "https://github.com/Echo-Smith", 7.0)

    add_section_heading(doc, "个人概述")
    add_body(doc, "具备近 2 年城市品牌传播、内容平台与跨团队项目推进经验，习惯从用户任务和业务流程中识别问题，明确优先级、版本边界与验收标准。个人发起 AI 内容产品“笔润智谈”，独立完成需求定义、版本规划、交互设计、Web 交付和用户迭代；V1 覆盖 200+ 名体验用户，其中 53 名作者完成真实写作任务。能够根据任务风险划分模型、规则、工具与人工确认的边界，并通过用户反馈、用例评测和运行数据持续调整方案。", size=8.2, after=1.4)

    add_section_heading(doc, "核心能力")
    add_compact_bullet(doc, bullet_id, "产品规划与体验设计：", "用户任务与流程洞察、问题拆解、MVP / 版本规划、优先级判断、PRD / 交互原型与上线验收。", 7.9)
    add_compact_bullet(doc, bullet_id, "AI 产品机制设计：", "根据任务不确定性与错误成本划分模型、规则、工具和人工边界，将 Agent Workflow、RAG、Memory、MCP 与 Function Calling 转化为可解释、可回退、可人工接管的产品机制。", 7.9)
    add_compact_bullet(doc, bullet_id, "验证与交付：", "用户反馈、用例与模型评测、Badcase 复盘、运行数据分析；可借助 AI Coding、Python、SQL 完成原型验证与数据整理。", 7.9)

    add_section_heading(doc, "核心项目")
    add_item_heading(doc, "笔润智谈｜面向作者与编辑协作的 AI 写作产品", "2026.03—至今")
    add_meta_links(doc, "个人发起｜负责需求定义、版本规划、交互设计、Agent 工作流与 Web 交付｜V1 真实使用，V2 内测", [("V1", "https://luminbuddy.ericdocmic.top"), ("V2", "https://luminbuddy2.ericdocmic.top"), ("GitHub", "https://github.com/Echo-Smith/luminbuddy-writing-agent-v2")])
    add_compact_bullet(doc, bullet_id, "定义问题与版本边界：", "从作者写作、编辑反馈和集中查重任务中，将问题拆为资料核验、表达同质化、个人偏好与返工协作；V1 先验证写作、润色和查重，V2 再扩展为研究、写作、审校的版本化协作流。", 8.0)
    add_compact_bullet(doc, bullet_id, "设计可控的人机协作：", "为兼顾关键步骤的确定性与复杂任务的灵活性，以固定 Pipeline 保证流程顺序，由 Agent 按上下文调用工具；根据任务质量决定继续流转、有限重试或转交用户批准、退回与接管。", 8.0)
    add_compact_bullet(doc, bullet_id, "依据反馈调整机制：", "针对“逻辑增强但格式趋同”、标题伦理和查重失败等反馈，分别调整写作规范、个性化记忆与并发兜底。V1 覆盖 200+ 名体验用户，53 名作者完成真实写作任务；批量链路完成 30 余次真实任务，单次处理 50+ 人 × 每人 30 条文本并在 5 分钟内全部返回；V2 已沉淀近 70 条完整样本，仍在验证稳定性与产品效果。", 8.0)

    add_item_heading(doc, "HearHer 听见妈妈｜面向妈妈情绪场景的安全音乐陪伴产品", "2026.06—2026.07")
    add_meta_links(doc, "腾讯音乐 AI Hackathon 两人团队｜产品规划、AI 能力设计、模型评测与 PWA 交付", [("项目介绍", "https://hearher.me"), ("项目页", "https://ideas.qq.com/work/detail/nBwZmb2DYJ1Ly3FqYh"), ("GitHub", "https://github.com/Echo-Smith/HearHer-PWA")])
    add_compact_bullet(doc, bullet_id, "", "围绕“表达当前状态并获得合适音乐陪伴”的任务，设计“状态输入—风险判断—受控推荐—真实播放—反馈沉淀”流程；由 LLM 负责语义理解和陪伴表达，以安全门槛、受控曲库与规则降级约束输出范围。", 8.0)
    add_compact_bullet(doc, bullet_id, "", "在确定方案前，以 177 条音频和 30 条产品用例对比 4 款 ASR、4 款 LLM，围绕识别效果、指令遵循与场景适配完成 700+ 次调用，为模型选择和交互方案提供依据；项目从 200+ 份提案中入围 14 项决赛作品。", 8.0)

    add_item_heading(doc, "贝伴｜订阅管理场景的 AI 操作助手原型")
    add_meta_links(doc, "自然语言管理订阅、查询支出与支付截图识别；结构化执行关键操作，对删除、金额修改和低置信结果保留确认", [("在线体验", "https://subbuddy.ericdocmic.top")])

    add_section_heading(doc, "工作经历")
    add_item_heading(doc, "杭州都市快报研学文化传播有限公司｜内容运营·城市品牌传播与内容策略", "2024.11—至今")
    add_compact_bullet(doc, bullet_id, "城市品牌内容体系：", "围绕城市整体形象，将节庆节点、人文经济与日常议题转化为持续内容供给，串联选题、约稿、审核、平台适配、发布监测与归档；作为协同中间层，2025 年推动 99 件团队稿件获相关专栏刊发。", 8.0)
    add_compact_bullet(doc, bullet_id, "多平台分发策略：", "承担今日头条、微博、小红书等渠道分发，根据平台内容形态、传播节点与用户反馈调整选题呈现、发布时间和更新节奏，日均推送重点内容约 3 条；2025 年今日头条日均展现约 3 万，粉丝突破 2 万。", 8.0)
    add_compact_bullet(doc, bullet_id, "传播验证与复盘：", "2026 年上半年深度参与 26 次模拟演练的出题 / 组织、数据分析与赛后复盘，识别协作和分发环节的薄弱点，支持区域综合排名保持杭州市前列。", 8.0)
    add_compact_bullet(doc, bullet_id, "业务系统验收：", "从一线使用流程出发参与内部系统上线前验收，围绕任务是否可完成、信息是否完整和反馈是否清晰提交 6 项分级问题；整理复现路径、关键字段与截图，帮助开发团队约 30 分钟定位关键问题并完成验证。", 8.0)

    add_section_heading(doc, "教育背景")
    add_item_heading(doc, "齐鲁工业大学｜新媒体技术（工科·计算机大类）｜本科", "2020.09—2024.06")
    add_meta(doc, "相关课程：机器学习与模式识别、人工智能与深度学习、智能媒体传播、移动软件开发")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_rhythm(footer, after=0, line=1.0)
    add_page_fields(footer)

    props = doc.core_properties
    props.title = "莫思意 - AI 产品经理定向简历"
    props.subject = "AI 应用 / Agent 工作流"
    props.author = "莫思意"
    props.keywords = "AI 产品经理, Agent, Workflow, RAG, MCP, Function Calling"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
