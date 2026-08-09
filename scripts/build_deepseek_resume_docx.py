from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docx" / "莫思意-DeepSeek-AI产品经理-定向简历.docx"
AVATAR = ROOT / "output" / "html" / "个人头像.jpg"

INK = "162235"
BODY = "34445A"
MUTED = "66758A"
ACCENT = "D1E0F9"
ACCENT_SOFT = "F2F6FD"
ACCENT_DEEP = "315F94"
# LibreOffice can ignore the East Asia font mapping when the same run mixes
# Chinese and Latin characters. Use one Unicode family for every script so the
# exported PDF keeps CJK glyphs while remaining editable in Word on macOS.
FONT_LATIN = "Arial Unicode MS"
FONT_CJK = "Arial Unicode MS"


def set_cell_free_font(run, size, color=BODY, bold=False, italic=False):
    run.font.name = FONT_LATIN
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), FONT_LATIN)
    fonts.set(qn("w:hAnsi"), FONT_LATIN)
    fonts.set(qn("w:eastAsia"), FONT_CJK)


def set_paragraph_rhythm(paragraph, before=0, after=2.6, line=1.14, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    fmt.line_spacing = line
    fmt.keep_together = keep


def shade_paragraph(paragraph, fill):
    ppr = paragraph._p.get_or_add_pPr()
    shd = ppr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        ppr.append(shd)
    shd.set(qn("w:fill"), fill)


def left_border(paragraph, color=ACCENT_DEEP, size="18", space="4"):
    ppr = paragraph._p.get_or_add_pPr()
    pbdr = ppr.find(qn("w:pBdr"))
    if pbdr is None:
        pbdr = OxmlElement("w:pBdr")
        ppr.append(pbdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), space)
    left.set(qn("w:color"), color)
    pbdr.append(left)


def add_hyperlink(paragraph, text, url, size=7.6):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT_DEEP)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), FONT_LATIN)
    rfonts.set(qn("w:hAnsi"), FONT_LATIN)
    rfonts.set(qn("w:eastAsia"), FONT_CJK)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    rpr.extend([rfonts, color, underline, sz])
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.extend([rpr, text_element])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(9.4)
    normal.font.color.rgb = RGBColor.from_string(BODY)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.paragraph_format.space_after = Pt(2.6)
    normal.paragraph_format.line_spacing = 1.14

    for name, size, color, before, after in [
        ("Heading 1", 11.0, INK, 7, 4.2),
        ("Heading 2", 10.1, INK, 4.8, 1.5),
        ("Heading 3", 9.5, INK, 3, 1.2),
    ]:
        style = styles[name]
        style.font.name = FONT_LATIN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Resume Meta" not in [s.name for s in styles]:
        meta = styles.add_style("Resume Meta", WD_STYLE_TYPE.PARAGRAPH)
    else:
        meta = styles["Resume Meta"]
    meta.font.name = FONT_LATIN
    meta.font.size = Pt(7.5)
    meta.font.color.rgb = RGBColor.from_string(MUTED)
    meta._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    meta.paragraph_format.space_after = Pt(1.4)
    meta.paragraph_format.line_spacing = 1.05


def add_custom_bullet_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•")
    lvl_jc = OxmlElement("w:lvlJc")
    lvl_jc.set(qn("w:val"), "left")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "240")
    tabs.append(tab)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "240")
    ind.set(qn("w:hanging"), "160")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "10")
    spacing.set(qn("w:line"), "250")
    spacing.set(qn("w:lineRule"), "auto")
    ppr.extend([tabs, ind, spacing])
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), ACCENT_DEEP)
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rpr.extend([rfonts, color])
    level.extend([start, num_fmt, lvl_text, lvl_jc, ppr, rpr])
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_bullet(paragraph, num_id):
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_el = OxmlElement("w:numId")
    num_id_el.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_el])
    ppr.append(num_pr)


def add_section_heading(doc, text):
    p = doc.add_paragraph(style="Heading 1")
    set_paragraph_rhythm(p, before=6.5, after=4.0, line=1.0, keep=True)
    shade_paragraph(p, ACCENT)
    left_border(p)
    p.paragraph_format.left_indent = Mm(2.0)
    p.paragraph_format.right_indent = Mm(1.5)
    r = p.add_run(text)
    set_cell_free_font(r, 11.0, INK, bold=True)
    return p


def add_item_heading(doc, title, date=None):
    p = doc.add_paragraph(style="Heading 2")
    set_paragraph_rhythm(p, before=4.2, after=1.1, line=1.0, keep=True)
    p.paragraph_format.tab_stops.add_tab_stop(Mm(186), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run(title)
    set_cell_free_font(r, 10.1, INK, bold=True)
    if date:
        p.add_run("\t")
        d = p.add_run(date)
        set_cell_free_font(d, 7.2, MUTED)
    return p


def add_meta(doc, text):
    p = doc.add_paragraph(style="Resume Meta")
    set_paragraph_rhythm(p, after=1.3, line=1.04, keep=True)
    r = p.add_run(text)
    set_cell_free_font(r, 7.5, MUTED)
    return p


def add_links(doc, links):
    p = doc.add_paragraph(style="Resume Meta")
    set_paragraph_rhythm(p, after=2.0, line=1.0, keep=True)
    # Keep a project's link row attached to the first evidence paragraph so a
    # page cannot end with an orphaned project title/link strip.
    p.paragraph_format.keep_with_next = True
    for index, (label, url) in enumerate(links):
        if index:
            sep = p.add_run("  |  ")
            set_cell_free_font(sep, 7.3, MUTED)
        add_hyperlink(p, label, url, 7.4)
    return p


def add_bullet(doc, num_id, label, text, size=9.1):
    p = doc.add_paragraph()
    apply_bullet(p, num_id)
    set_paragraph_rhythm(p, after=2.1, line=1.12, keep=True)
    if label:
        lead = p.add_run(label)
        set_cell_free_font(lead, size, ACCENT_DEEP, bold=True)
    body = p.add_run(text)
    set_cell_free_font(body, size, BODY)
    return p


def add_body(doc, text, size=9.35, after=2.6, fill=None):
    p = doc.add_paragraph()
    set_paragraph_rhythm(p, after=after, line=1.14)
    if fill:
        shade_paragraph(p, fill)
        p.paragraph_format.left_indent = Mm(2)
        p.paragraph_format.right_indent = Mm(2)
    r = p.add_run(text)
    set_cell_free_font(r, size, BODY)
    return p


def add_page_fields(paragraph):
    paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(186), WD_TAB_ALIGNMENT.RIGHT)
    left = paragraph.add_run("莫思意 | DeepSeek AI 产品经理")
    set_cell_free_font(left, 7.0, MUTED)
    paragraph.add_run("\t")
    for field_name in ["PAGE", "NUMPAGES"]:
        if field_name == "NUMPAGES":
            slash = paragraph.add_run(" / ")
            set_cell_free_font(slash, 7.0, MUTED)
        fld = OxmlElement("w:fldSimple")
        fld.set(qn("w:instr"), field_name)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), MUTED)
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), "14")
        rpr.extend([color, sz])
        text = OxmlElement("w:t")
        text.text = "1"
        run.extend([rpr, text])
        fld.append(run)
        paragraph._p.append(fld)


def add_first_page_avatar(section):
    """Place the headshot in the first-page header without disturbing ATS flow."""
    section.different_first_page_header_footer = True
    header = section.first_page_header
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run()
    inline_shape = run.add_picture(str(AVATAR), width=Mm(20.5))
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("name", "莫思意头像")
    doc_pr.set("descr", "莫思意证件照")

    # Convert the inline drawing to a page-positioned anchor. This keeps the
    # portrait in the visual header while leaving the resume's text order and
    # first-page vertical budget unchanged.
    anchor = inline_shape._inline
    anchor.tag = qn("wp:anchor")
    for key, value in {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "251658240",
        "behindDoc": "0",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }.items():
        anchor.set(key, value)

    simple_pos = OxmlElement("wp:simplePos")
    simple_pos.set("x", "0")
    simple_pos.set("y", "0")

    position_h = OxmlElement("wp:positionH")
    position_h.set("relativeFrom", "margin")
    h_align = OxmlElement("wp:align")
    h_align.text = "right"
    position_h.append(h_align)

    position_v = OxmlElement("wp:positionV")
    position_v.set("relativeFrom", "margin")
    v_offset = OxmlElement("wp:posOffset")
    v_offset.text = "0"
    position_v.append(v_offset)

    extent = anchor.find(qn("wp:extent"))
    extent_index = anchor.index(extent)
    anchor.insert(0, simple_pos)
    anchor.insert(1, position_h)
    anchor.insert(2, position_v)

    wrap_none = OxmlElement("wp:wrapNone")
    anchor.insert(extent_index + 4, wrap_none)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(10)
    section.bottom_margin = Mm(9)
    section.left_margin = Mm(12)
    section.right_margin = Mm(12)
    section.header_distance = Mm(5)
    section.footer_distance = Mm(5)

    add_first_page_avatar(section)

    configure_styles(doc)
    bullet_id = add_custom_bullet_numbering(doc)

    title = doc.add_paragraph()
    set_paragraph_rhythm(title, after=1.0, line=1.0)
    r = title.add_run("莫思意")
    set_cell_free_font(r, 22.0, INK, bold=True)

    target = doc.add_paragraph()
    set_paragraph_rhythm(target, after=1.5, line=1.0)
    r = target.add_run("AI 产品经理｜AI 产品 / Agent Harness")
    set_cell_free_font(r, 10.3, ACCENT_DEEP, bold=True)

    contact = doc.add_paragraph()
    set_paragraph_rhythm(contact, after=4.0, line=1.0)
    for text in ["19290865729", "aodomosiyi@foxmail.com", "杭州"]:
        rr = contact.add_run(text + "  |  ")
        set_cell_free_font(rr, 7.6, MUTED)
    add_hyperlink(contact, "作品集", "https://profile.ericdocmic.top", 7.6)
    sep = contact.add_run("  |  ")
    set_cell_free_font(sep, 7.6, MUTED)
    add_hyperlink(contact, "GitHub", "https://github.com/Echo-Smith", 7.6)

    add_section_heading(doc, "个人概述")
    add_body(
        doc,
        "2024 届新媒体技术本科，现工作于内容生产与创作者协作一线；个人发起并独立构建 AI 内容产品“笔润智谈”，完成用户问题发现、版本规划、产品设计、Agent Harness 实现、Web 交付与反馈迭代。V1 覆盖 200+ 名体验用户，其中 53 名作者完成真实写作任务；V2 截至 2026.07.31 有 14 名用户、近 70 条完整任务样本。实际接入 DeepSeek API，长期使用 Codex、Claude Code、Dify 等 Agent / AI Coding / Workflow 工具推进调研、原型、开发、测试与复盘。",
        size=9.35,
        after=2.4,
    )
    p = doc.add_paragraph()
    set_paragraph_rhythm(p, after=3.0, line=1.12)
    shade_paragraph(p, ACCENT_SOFT)
    p.paragraph_format.left_indent = Mm(2)
    p.paragraph_format.right_indent = Mm(2)
    lead = p.add_run("产品判断：")
    set_cell_free_font(lead, 9.2, ACCENT_DEEP, bold=True)
    body = p.add_run("Agent 的价值不在最大化自主性，而在真实任务中把 Context、Tools、Memory、退出边界与评测组织成可观察、可干预、可持续迭代的运行环境；高置信步骤交给模型，关键决策保留依据、撤销和人工接管。")
    set_cell_free_font(body, 9.2, BODY)

    add_section_heading(doc, "核心能力")
    add_bullet(doc, bullet_id, "产品与用户：", "用户反馈与定性观察、需求拆解、MVP / 版本规划、PRD、优先级判断、交互原型、上线验收、Badcase 与版本复盘。", 8.95)
    add_bullet(doc, bullet_id, "Agent 产品：", "LLM API、Agent Loop、Tool Use、Reasoning / Planning、Prompt / Context Engineering、RAG、Function Calling、Memory、MCP 原型、Multi-Agent、Human-in-the-loop。", 8.95)
    add_bullet(doc, bullet_id, "数据与工程：", "用例集与模型评测、Trace / Token / 时延 / 错误分析；Python、SQL / PostgreSQL、Go、TypeScript / React、REST API、WebSocket、Docker、Git，可独立交付 Web / PWA 产品原型。", 8.95)

    add_section_heading(doc, "核心项目")
    add_item_heading(doc, "笔润智谈｜面向真实内容生产的 Chat / Agent 工作台", "2026.03—至今")
    add_meta(doc, "个人发起并独立完成产品规划、交互设计、Agent 应用链路与 Web 交付｜V1 已进入内部内容生产流程，V2 持续内测")
    add_links(doc, [
        ("V1", "https://luminbuddy.ericdocmic.top"),
        ("V2", "https://luminbuddy2.ericdocmic.top"),
        ("GitHub", "https://github.com/Echo-Smith/luminbuddy-writing-agent-v2"),
    ])
    add_bullet(doc, bullet_id, "从真实任务规划版本：", "在作者、编辑与内容运营流程中识别资料核验、表达同质化、个人视角、版本追踪和返工协作问题；先用 V1 验证写作、润色与批量查重，再将 V2 定义为可追踪、可回退、可人工接管的 Agent 编辑部。")
    add_bullet(doc, bullet_id, "让反馈进入产品机制：", "根据“逻辑更强但格式趋同”的反馈调整写作规范、减少模板化句式；面对标题伦理与个人视角分歧，将意见建模为按用户、来源和范围生效的个性化记忆，避免把一次反馈固化为全局 Prompt。")
    add_bullet(doc, bullet_id, "平衡确定性与自主性：", "保留固定 Pipeline 保障关键步骤顺序，以 ReAct UnifiedAgent 按上下文选择工具；设置取消、全局超时、Token 预算、断连暂停和连续模型失败断路器，并以迭代上限与人工确认超时避免任务失控。")
    add_bullet(doc, bullet_id, "设计 Context、Memory 与 Tool 边界：", "以工作摘要、会话消息、长期偏好和实体关系网络构成四层记忆，提供 Markdown 与数据库导入 / 导出；实现 MCP 客户端接入与进程内 Server 原型，将外部工具和内置能力注册进统一 ToolRegistry。")
    add_bullet(doc, bullet_id, "编排 Agent 协作与人工决策：", "将研究、写作、审校三类 Agent 组织为版本化协作流，以数据库 Lease 防止重复执行，按信源、信息缺口、稿件结构与审校严重度动态路由；低质量任务有限重试后进入人工 Decision，支持批准、退回和接管。")
    add_bullet(doc, bullet_id, "用数据验证与复盘：", "接入 DeepSeek API 处理生成、审校和记忆抽取，以用例、Badcase、Trace、Token、时延、错误与人工介入记录复盘链路。V1 覆盖 200+ 名体验用户，其中 53 名作者完成真实任务；V2 有 14 名用户、近 70 条完整任务样本。重构并发查重与失败兜底后完成 30 余次任务，单次处理 50 余人 × 每人 30 条文本，5 分钟内全部返回。")

    add_item_heading(doc, "HearHer 听见妈妈｜多模态体验、模型评测与安全边界", "2026.06—2026.07")
    add_meta(doc, "腾讯音乐 AI Hackathon 两人团队｜负责产品规划、AI 能力设计、模型评测与 PWA 交付")
    add_links(doc, [
        ("项目介绍", "https://hearher.me"),
        ("黑客松项目页", "https://ideas.qq.com/work/detail/nBwZmb2DYJ1Ly3FqYh"),
        ("GitHub", "https://github.com/Echo-Smith/HearHer-PWA"),
    ])
    add_bullet(doc, bullet_id, "", "面向妈妈群体的情绪与音乐需求，设计“状态输入—安全判断—受控推荐—真实播放—反馈沉淀”流程；由 LLM 负责语义理解和陪伴表达，以安全硬门槛、受控曲库与规则降级约束输出范围。")
    add_bullet(doc, bullet_id, "", "使用 177 条音频与 30 条产品用例对比 4 款 ASR、4 款 LLM 的识别、指令遵循和产品表现，完成 700+ 次调用；项目从 200+ 份提案中入围 14 项决赛作品。")

    add_item_heading(doc, "贝伴｜AI 订阅管理助手")
    add_links(doc, [("在线体验", "https://subbuddy.ericdocmic.top")])
    add_body(doc, "设计 Function Calling 协议处理订阅增删改查与支出查询；视觉模型提取支付截图字段，并对删除、金额修改和低置信结果保留人工确认，形成“模型建议—用户确认—系统执行”的安全交互。", size=9.2, after=3.0)

    add_section_heading(doc, "工作经历")
    add_item_heading(doc, "杭州都市快报研学文化传播有限公司｜内容运营", "2024.11—至今")
    add_bullet(doc, bullet_id, "", "负责小红书、今日头条、微博等 40 万粉级官方账号的选题、生产、分发和复盘，协调 30+ 本地 KOL / KOC，参与约 150 人素人号矩阵管理；持续从作者、编辑和运营反馈中识别内容约束与协作问题。")
    add_bullet(doc, bullet_id, "", "使用飞书智能表格搭建话题获取、分类与数据看板，年度汇总 75 个话题、累计阅读量约 46.3 亿；使用 Python 清洗多平台数据，支持选题判断和阶段复盘。")
    add_bullet(doc, bullet_id, "", "参与内部 SaaS 上线前验收，围绕业务流程、数据字段和交互体验提交 6 项分级问题，整理复现路径、关键字段与截图，协助开发团队在半小时内定位关键问题并完成核心功能验证。")

    add_section_heading(doc, "教育背景")
    add_item_heading(doc, "齐鲁工业大学｜新媒体技术（工科·计算机大类）｜本科", "2020.09—2024.06")
    add_meta(doc, "相关课程：机器学习与模式识别、人工智能与深度学习、智能媒体传播、移动软件开发")
    add_meta(doc, "证书：大学英语六级 425 分、普通话水平测试二级甲等")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_rhythm(footer, after=0, line=1.0)
    add_page_fields(footer)

    props = doc.core_properties
    props.title = "莫思意 - DeepSeek AI 产品经理定向简历"
    props.subject = "AI 产品 / Agent Harness"
    props.author = "莫思意"
    props.keywords = "AI 产品经理, Agent Harness, DeepSeek, Agent, Memory, MCP"

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
