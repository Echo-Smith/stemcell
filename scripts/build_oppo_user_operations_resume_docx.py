from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from build_deepseek_resume_docx import (
    ACCENT_DEEP,
    BODY,
    INK,
    MUTED,
    add_body,
    add_custom_bullet_numbering,
    add_first_page_avatar,
    add_hyperlink,
    add_item_heading,
    add_meta,
    add_section_heading,
    apply_bullet,
    configure_styles,
    set_cell_free_font,
    set_paragraph_rhythm,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "output" / "docx" / "莫思意-OPPO-高级用户运营经理-转向简历.docx"


def add_page_fields(paragraph):
    paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(186), WD_TAB_ALIGNMENT.RIGHT)
    left = paragraph.add_run("莫思意 | 用户运营 / 产品运营")
    set_cell_free_font(left, 7.0, MUTED)
    paragraph.add_run("\t")
    for field_name in ["PAGE", "NUMPAGES"]:
        if field_name == "NUMPAGES":
            set_cell_free_font(paragraph.add_run(" / "), 7.0, MUTED)
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), field_name)
        run = OxmlElement("w:r")
        prop = OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), MUTED)
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "14")
        prop.extend([color, size])
        text = OxmlElement("w:t")
        text.text = "1"
        run.extend([prop, text])
        field.append(run)
        paragraph._p.append(field)


def add_compact_bullet(doc, num_id, label, text, size=8.45):
    paragraph = doc.add_paragraph()
    apply_bullet(paragraph, num_id)
    set_paragraph_rhythm(paragraph, after=0.75, line=1.055, keep=True)
    if label:
        set_cell_free_font(paragraph.add_run(label), size, ACCENT_DEEP, bold=True)
    set_cell_free_font(paragraph.add_run(text), size, BODY)
    return paragraph


def add_meta_links(doc, text, links):
    paragraph = doc.add_paragraph(style="Resume Meta")
    set_paragraph_rhythm(paragraph, after=0.75, line=1.0, keep=True)
    paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(190), WD_TAB_ALIGNMENT.RIGHT)
    set_cell_free_font(paragraph.add_run(text), 6.9, MUTED)
    paragraph.add_run("\t")
    for index, (label, url) in enumerate(links):
        if index:
            set_cell_free_font(paragraph.add_run(" | "), 6.7, MUTED)
        add_hyperlink(paragraph, label, url, 6.8)
    return paragraph


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(7)
    section.bottom_margin = Mm(6)
    section.left_margin = Mm(9)
    section.right_margin = Mm(9)
    section.header_distance = Mm(3)
    section.footer_distance = Mm(3.5)
    add_first_page_avatar(section)

    configure_styles(doc)
    doc.styles["Normal"].font.size = Pt(8.55)
    doc.styles["Normal"].paragraph_format.space_after = Pt(1.45)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.055
    doc.styles["Heading 1"].font.size = Pt(10.0)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(3.2)
    doc.styles["Heading 1"].paragraph_format.space_after = Pt(2.2)
    doc.styles["Heading 2"].font.size = Pt(9.2)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(2.2)
    doc.styles["Heading 2"].paragraph_format.space_after = Pt(0.7)
    bullet_id = add_custom_bullet_numbering(doc)

    title = doc.add_paragraph()
    set_paragraph_rhythm(title, after=1.0, line=1.0)
    set_cell_free_font(title.add_run("莫思意"), 19.5, INK, bold=True)

    target = doc.add_paragraph()
    set_paragraph_rhythm(target, after=1.5, line=1.0)
    set_cell_free_font(target.add_run("用户运营 / 产品运营｜创作者生态与内容增长"), 9.4, ACCENT_DEEP, bold=True)

    contact = doc.add_paragraph()
    set_paragraph_rhythm(contact, after=2.0, line=1.0)
    for value in ["19290865729", "aodomosiyi@foxmail.com", "杭州"]:
        set_cell_free_font(contact.add_run(value + "  |  "), 7.0, MUTED)
    add_hyperlink(contact, "作品集", "https://profile.ericdocmic.top", 7.0)
    set_cell_free_font(contact.add_run("  |  "), 7.0, MUTED)
    add_hyperlink(contact, "GitHub", "https://github.com/Echo-Smith", 7.0)

    add_section_heading(doc, "个人概述")
    add_body(doc, "具备近 2 年内容平台、创作者协作与产品化实践，服务 40 万粉级官方账号，协调 30+ 本地 KOL/KOC，参与约 150 人素人号矩阵运营，覆盖创作者沟通、内容供给、平台分发、数据监测与反馈复盘。能独立完成视频拍摄、剪辑与 AIGC 影像内容；个人发起“笔润智谈”，将 53 名作者的真实任务与反馈转化为产品迭代。", size=8.55, after=1.45)

    add_section_heading(doc, "核心能力")
    add_compact_bullet(doc, bullet_id, "创作者与用户运营：", "KOL/KOC 协作、账号矩阵、任务分发、发布巡检、反馈回收与核心用户沟通。", 8.25)
    add_compact_bullet(doc, bullet_id, "内容与活动运营：", "节庆 / 人文经济 / 城市话题策划，达人对接、内容审核、现场执行、多平台分发和效果复盘。", 8.25)
    add_compact_bullet(doc, bullet_id, "用户洞察与产品优化：", "将用户原话拆为共性问题、个性偏好和流程故障，结合优先级、验收标准和数据复盘推动迭代。", 8.25)
    add_compact_bullet(doc, bullet_id, "影像内容与数据分析：", "短视频拍摄剪辑、AIGC 图像 / 视频工作流；Excel、飞书智能表格、Python 数据清洗与看板整理。", 8.25)

    add_section_heading(doc, "工作经历")
    add_item_heading(doc, "杭州都市快报研学文化传播有限公司｜内容运营·城市品牌传播与内容策略", "2024.11—至今")
    add_compact_bullet(doc, bullet_id, "创作者与矩阵协作：", "服务 40 万粉级官方账号，协调 30+ 本地 KOL/KOC，参与约 150 人素人号矩阵的账号激活、任务分发、发布巡检和反馈回收；根据账号定位与平台规则协同内容交付。", 8.45)
    add_compact_bullet(doc, bullet_id, "内容供给与活动执行：", "围绕节庆、人文经济与城市话题，串联选题、约稿、审核优化、平台适配、发布监测与归档；参与达人对接、内容包装、现场带队和多方协作，2025 年协同推进 99 件团队稿件获相关专栏刊发。", 8.45)
    add_compact_bullet(doc, bullet_id, "平台策略与数据复盘：", "负责小红书、今日头条、微博等渠道的选题、内容适配、发布排期和效果监测；代表账号 7 天发布 7 篇，获得 10.4 万观看、5760 次互动，观看与互动均超过 95% 同类作者；代表话题进入微博杭州同城榜 TOP1，单专题阅读或播放超 10 万。", 8.45)
    add_compact_bullet(doc, bullet_id, "用户支持与体验优化：", "参与 App 提醒推送、用户答疑、扩容登记及在线训练营任务分发和账号维护监督；从一线使用流程出发参与内部 SaaS 上线前验收，提交 6 项分级问题，帮助开发团队约 30 分钟定位关键问题并完成验证。", 8.45)

    add_section_heading(doc, "代表实践")
    add_item_heading(doc, "笔润智谈｜创作者共创与内容工具迭代", "2026.03—至今")
    add_meta_links(doc, "个人发起｜需求梳理、产品设计、用户反馈与版本迭代", [("V1", "https://luminbuddy.ericdocmic.top"), ("V2", "https://luminbuddy2.ericdocmic.top")])
    add_compact_bullet(doc, bullet_id, "用户与任务：", "从作者写作、编辑反馈和集中查重中识别资料分散、表达同质化、个人偏好与返工问题；V1 覆盖 200+ 名体验用户，其中 53 名作者完成真实写作任务。", 8.45)
    add_compact_bullet(doc, bullet_id, "共创与迭代：", "针对“逻辑更强但格式趋同”、标题与个人视角冲突、批量任务失败等反馈，分别调整写作规范、按用户保存偏好并重构失败兜底；把单次意见分成全局规则、个性约束与流程问题。", 8.45)
    add_compact_bullet(doc, bullet_id, "任务交付：", "完成 30 余次批量真实任务，单次处理 50+ 人 × 每人 30 条文本，并在 5 分钟内全部返回；V2 持续记录任务状态、用户反馈与完整样本，仍在验证稳定性与产品效果。", 8.45)

    add_item_heading(doc, "影像与 AIGC 内容创作实践")
    add_compact_bullet(doc, bullet_id, "", "工作中独立完成短视频拍摄、剪辑与 AIGC 漫画创作；代表 AIGC 地标内容融合文生图、图生视频与非遗剪纸主题，单条播放超 10 万。", 8.45)
    add_compact_bullet(doc, bullet_id, "", "独立制作 21 页《AI 质感漫画与视频创作实操指南》，结合 5 个视频及多组生成案例，梳理从角色定妆、关键帧、动态生成到剪辑发布的工作流，并从人物一致性、时序连贯及光影、色彩、材质、细节、氛围等维度判断画面质量。", 8.45)

    add_section_heading(doc, "教育背景")
    add_item_heading(doc, "齐鲁工业大学｜新媒体技术（工科·计算机大类）｜本科", "2020.09—2024.06")
    add_meta(doc, "相关课程：智能媒体传播、移动软件开发、机器学习与模式识别、人工智能与深度学习")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_paragraph_rhythm(footer, after=0, line=1.0)
    add_page_fields(footer)

    props = doc.core_properties
    props.title = "莫思意 - OPPO 用户运营 / 产品运营转向简历"
    props.subject = "创作者生态与内容增长"
    props.author = "莫思意"
    props.keywords = "用户运营, 产品运营, 创作者生态, 内容增长, 影像创作"
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
